"""导出模块：txt / md / docx / srt。

负责人：gui-dev。接口契约（pipeline 依赖，改动需同步）：

    paths = export_result(result, out_dir=Path(...), base_name="会议记录",
                          formats=["txt", "md", "docx", "srt"],
                          title="会议记录", include_raw=True) -> list[Path]

要求：
* 有整理稿时，txt/md/docx 输出整理稿；同时另存一份
  ``<base_name>_原始转写.txt``（include_raw=True 时）。
* 没有整理稿时，txt/md/docx 输出原始转写。
* srt 始终由原始转写的时间戳生成。
* 文件名重复时自动加 ``(2)``、``(3)`` 后缀，绝不覆盖已有文件。
* 文件开头可附少量元信息（源文件、时长、生成时间、整理风格），
  保证正文仍可直接使用。
* 所有文本以 UTF-8 写入（docx 由 python-docx 生成）。

实现说明：
* 文件名会再次经过 :func:`voxscribe.util.sanitize_filename` 防御性清洗；
* 每个文件先确定唯一路径再写入，同一次导出的多个文件也不会互相覆盖；
* 任何写盘异常都包装为 :class:`voxscribe.errors.VoxError`（中文提示）；
* 正文按段落拆分写入 docx，markdown 风格的 ``#`` 标题与 ``-`` 列表会转为
  Word 标题/项目符号，便于阅读。
"""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from .errors import VoxError
from .models import EXPORT_FORMATS, JobResult, format_srt_timestamp
from .util import human_duration, sanitize_filename

log = logging.getLogger(__name__)

_VALID_FORMATS = tuple(EXPORT_FORMATS)
_DEFAULT_FORMAT = "txt"
_MAX_DUPLICATES = 9999

#: docx 中使用的字体（Windows 自带；python-docx 找不到字体时由 Word 回退）
_DOCX_FONT = "Microsoft YaHei"
_DOCX_TITLE_FONT = "Microsoft YaHei UI"


# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------
def _normalize_formats(formats: Optional[Sequence[str]]) -> List[str]:
    """清洗格式列表：去重、保序、忽略非法值；空列表回退到 txt。"""
    cleaned: List[str] = []
    for item in formats or ():
        if not isinstance(item, str):
            continue
        name = item.strip().lower().lstrip(".")
        if name in _VALID_FORMATS and name not in cleaned:
            cleaned.append(name)
    return cleaned or [_DEFAULT_FORMAT]


def _has_refined(result: JobResult) -> bool:
    return bool(result.refined and (result.refined.text or "").strip())


def _body_text(result: JobResult) -> str:
    """正文：整理稿优先，没有整理稿时退回原始转写。"""
    if _has_refined(result):
        assert result.refined is not None  # for type checkers
        return result.refined.text.strip()
    return result.transcript.text.strip()


def _raw_text(result: JobResult) -> str:
    """原始转写（优先带时间戳）。"""
    transcript = result.transcript
    stamped = transcript.with_timestamps().strip()
    return stamped or transcript.text.strip()


def _audio_duration(result: JobResult) -> float:
    """音频总时长：优先 transcript.duration，其次取最后一个片段的结束时间。"""
    duration = float(result.transcript.duration or 0.0)
    if duration <= 0:
        ends = [float(seg.end or 0.0) for seg in result.transcript.segments]
        duration = max(ends) if ends else 0.0
    return max(0.0, duration)


def _audio_name(result: JobResult) -> str:
    raw = (result.audio_path or "").strip()
    if not raw:
        return "（未知）"
    try:
        return Path(raw).name or raw
    except (OSError, ValueError):  # 极端非法路径
        return raw


def _metadata(
    result: JobResult, title: str
) -> List[Tuple[str, str]]:
    """构造元信息（键, 值）列表；顺序固定，便于文本/markdown/docx 共用。"""
    duration = _audio_duration(result)
    return [
        ("标题", title),
        ("源音频", _audio_name(result)),
        ("音频时长", human_duration(duration) if duration > 0 else "未知"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("内容", "AI 整理稿 · DeepSeek" if _has_refined(result) else "原始转写"),
    ]


def _build_txt(meta: List[Tuple[str, str]], body: str) -> str:
    header = "\n".join(f"{key}：{value}" for key, value in meta)
    return f"{header}\n\n{body}\n"


def _build_md(title: str, meta: List[Tuple[str, str]], body: str) -> str:
    quote = "\n".join(
        f"> {key}：{value}  " for key, value in meta if key != "标题"
    )
    return f"# {title}\n\n{quote}\n\n{body}\n"


def _build_raw_txt(result: JobResult, meta: List[Tuple[str, str]]) -> str:
    header = "\n".join(
        f"{key}：{value}" for key, value in meta if key not in ("标题", "内容")
    )
    header = f"{header}\n内容：原始转写（含时间戳）"
    return f"{header}\n{'=' * 40}\n{_raw_text(result)}\n"


def _split_srt_blocks(result: JobResult) -> List[str]:
    """按 transcript.segments 生成 SRT 块（空文本跳过，序号连续）。"""
    blocks: List[str] = []
    index = 0
    for segment in result.transcript.segments:
        text = (segment.text or "").strip()
        if not text:
            continue
        start = max(0.0, float(segment.start or 0.0))
        end = float(segment.end or 0.0)
        if end <= start:  # 防御异常时间段，保证字幕时长 > 0
            end = start + 0.5
        index += 1
        blocks.append(
            f"{index}\n"
            f"{format_srt_timestamp(start)} --> {format_srt_timestamp(end)}\n"
            f"{text}\n"
        )
    return blocks


def _build_srt(result: JobResult) -> str:
    return "\n".join(_split_srt_blocks(result))


# ---------------------------------------------------------------------------
# 路径与写入
# ---------------------------------------------------------------------------
def _ensure_out_dir(out_dir: Path) -> None:
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise VoxError(
            f"无法创建输出目录：{out_dir}（{exc}）",
            hint="请检查该路径是否有写入权限，或在「设置 → 默认输出目录」中换一个文件夹。",
        ) from exc
    if not out_dir.is_dir():
        raise VoxError(
            f"输出路径不是文件夹：{out_dir}",
            hint="请在「设置 → 默认输出目录」中选择一个文件夹。",
        )


def _unique_path(out_dir: Path, stem: str, suffix: str, reserved: set) -> Path:
    """返回不存在的路径；重名时依次尝试 `` (2)``、`` (3)``……"""
    candidate = out_dir / f"{stem}{suffix}"
    counter = 2
    while candidate.exists() or candidate in reserved:
        candidate = out_dir / f"{stem} ({counter}){suffix}"
        counter += 1
        if counter > _MAX_DUPLICATES:
            raise VoxError(
                "输出目录中同名文件过多，无法生成新的文件名",
                hint="请更换输出目录，或清理同名旧文件后重试。",
            )
    reserved.add(candidate)
    return candidate


def _write_text(path: Path, text: str) -> None:
    # 显式使用 LF，保证跨平台输出一致（Windows 记事本/Word 均可正常打开）
    path.write_text(text, encoding="utf-8", newline="\n")


def _add_docx_default_font(document, *, font: str, size_pt: float) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style = document.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font)


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")
_BULLET_RE = re.compile(r"^[-*+]\s+(\S.*)$")


def _write_docx(path: Path, title: str, meta: List[Tuple[str, str]], body: str) -> None:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except Exception as exc:  # pragma: no cover - 依赖缺失的兜底
        raise VoxError(
            "缺少 python-docx 依赖，无法导出 Word 文档",
            hint="请运行 setup.bat 安装依赖，或先只导出 txt / md 格式。",
        ) from exc

    document = Document()
    _add_docx_default_font(document, font=_DOCX_FONT, size_pt=11)
    try:
        title_style = document.styles["Title"]
        title_style.font.name = _DOCX_TITLE_FONT
        title_style.font.size = Pt(20)
    except KeyError:  # pragma: no cover - 模板差异
        pass

    document.add_heading(title, level=1)

    meta_paragraph = document.add_paragraph()
    meta_run = meta_paragraph.add_run(
        " ｜ ".join(f"{key}：{value}" for key, value in meta if key != "标题")
    )
    meta_run.italic = True
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    for line in body.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = _HEADING_RE.match(stripped)
        if heading:
            level = min(len(heading.group(1)) + 1, 4)
            document.add_heading(heading.group(2), level=level)
            continue
        bullet = _BULLET_RE.match(stripped)
        if bullet:
            try:
                document.add_paragraph(bullet.group(1), style="List Bullet")
            except KeyError:  # pragma: no cover - 模板差异
                document.add_paragraph(stripped)
            continue
        document.add_paragraph(stripped)

    document.save(str(path))


# ---------------------------------------------------------------------------
# 对外接口
# ---------------------------------------------------------------------------
def export_result(
    result: JobResult,
    *,
    out_dir: Path,
    base_name: str,
    formats: Sequence[str],
    title: Optional[str] = None,
    include_raw: bool = True,
) -> List[Path]:
    """按 ``formats`` 导出结果文件，返回实际写出的路径列表。

    :param result: 一个已完成的转写/整理结果。
    :param out_dir: 输出目录（不存在会自动创建）。
    :param base_name: 输出文件名（不含扩展名），函数内部会再次清洗。
    :param formats: txt / md / docx / srt 的任意组合；非法值忽略，空则至少 txt。
    :param title: 文档标题，默认使用 ``base_name``。
    :param include_raw: 有整理稿时是否额外写出 ``<base_name>_原始转写.txt``。
    """
    target_dir = Path(out_dir)
    _ensure_out_dir(target_dir)

    clean_base = sanitize_filename(
        base_name or Path(result.audio_path or "").stem or "转写结果"
    )
    doc_title = (title or clean_base).strip() or clean_base
    meta = _metadata(result, doc_title)
    body = _body_text(result)
    requested = _normalize_formats(formats)

    written: List[Path] = []
    reserved: set = set()

    for fmt in requested:
        filename_stem = clean_base
        if fmt == "txt":
            path = _unique_path(target_dir, filename_stem, ".txt", reserved)
            try:
                _write_text(path, _build_txt(meta, body))
            except VoxError:
                raise
            except Exception as exc:
                raise VoxError(
                    f"导出文本文件失败：{path.name}（{exc}）",
                    hint="请确认输出目录可写、磁盘空间充足，或更换输出目录后重试。",
                ) from exc
        elif fmt == "md":
            path = _unique_path(target_dir, filename_stem, ".md", reserved)
            try:
                _write_text(path, _build_md(doc_title, meta, body))
            except VoxError:
                raise
            except Exception as exc:
                raise VoxError(
                    f"导出 Markdown 文件失败：{path.name}（{exc}）",
                    hint="请确认输出目录可写、磁盘空间充足，或更换输出目录后重试。",
                ) from exc
        elif fmt == "docx":
            path = _unique_path(target_dir, filename_stem, ".docx", reserved)
            try:
                _write_docx(path, doc_title, meta, body)
            except VoxError:
                raise
            except Exception as exc:
                raise VoxError(
                    f"导出 Word 文档失败：{path.name}（{exc}）",
                    hint="请确认输出目录可写、磁盘空间充足，或先只导出 txt / md 格式。",
                ) from exc
        elif fmt == "srt":
            blocks = _split_srt_blocks(result)
            if not blocks:
                log.warning("转写结果没有可生成字幕的片段，跳过 srt 导出")
                continue
            path = _unique_path(target_dir, filename_stem, ".srt", reserved)
            try:
                _write_text(path, "\n".join(blocks) + "\n")
            except VoxError:
                raise
            except Exception as exc:
                raise VoxError(
                    f"导出字幕文件失败：{path.name}（{exc}）",
                    hint="请确认输出目录可写、磁盘空间充足，或更换输出目录后重试。",
                ) from exc
        else:  # pragma: no cover - _normalize_formats 已过滤
            continue
        written.append(path)
        log.info("已导出 %s", path)

    # 有整理稿时另存原始转写，便于校对
    if include_raw and _has_refined(result):
        raw_path = _unique_path(
            target_dir, f"{clean_base}_原始转写", ".txt", reserved
        )
        try:
            _write_text(raw_path, _build_raw_txt(result, meta))
        except VoxError:
            raise
        except Exception as exc:
            raise VoxError(
                f"导出原始转写失败：{raw_path.name}（{exc}）",
                hint="请确认输出目录可写、磁盘空间充足，或更换输出目录后重试。",
            ) from exc
        written.append(raw_path)
        log.info("已导出原始转写 %s", raw_path)

    return written


__all__ = ["export_result"]
