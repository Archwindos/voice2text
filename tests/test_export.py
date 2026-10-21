"""导出模块测试：txt / md / docx / srt、元信息、重名递增、中文与长文本。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from voxscribe.errors import VoxError
from voxscribe.export import export_result
from voxscribe.models import JobResult, RefineResult, Segment, Transcript

REFINED_TEXT = (
    "今天主要讨论项目进度问题。\n\n"
    "前端大约完成了百分之七十。\n"
    "后端还需要两周时间。"
)


def _make_result(
    transcript: Transcript,
    refined_text: str | None = REFINED_TEXT,
    audio_name: str = "会议 录音.mp3",
) -> JobResult:
    refined = RefineResult(text=refined_text) if refined_text is not None else None
    return JobResult(
        audio_path=str(Path("D:/示例音频") / audio_name),
        transcript=transcript,
        refined=refined,
        elapsed=12.5,
    )


def _names(paths: list[Path]) -> list[str]:
    return [path.name for path in paths]


# ---------------------------------------------------------------------------
# txt / md / docx 基本导出
# ---------------------------------------------------------------------------
def test_export_txt_with_refined_and_raw(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    out_dir = tmp_path / "out"
    files = export_result(
        result,
        out_dir=out_dir,
        base_name="会议记录",
        formats=["txt"],
        title="周会纪要",
        include_raw=True,
    )

    assert _names(files) == ["会议记录.txt", "会议记录_原始转写.txt"]
    assert all(path.exists() for path in files)

    main_text = files[0].read_text(encoding="utf-8")
    # 元信息
    assert main_text.startswith("标题：周会纪要")
    assert "源音频：会议 录音.mp3" in main_text
    assert "音频时长：10秒" in main_text  # duration=9.5 -> 四舍五入 10 秒
    assert "生成时间：" in main_text
    assert "内容：AI 整理稿" in main_text
    # 元信息与正文之间空行分隔
    header, body = main_text.split("\n\n", 1)
    assert "标题：" in header
    # 正文使用整理稿
    assert "前端大约完成了百分之七十" in body
    assert "嗯，那个" not in body

    raw_text = files[1].read_text(encoding="utf-8")
    assert "内容：原始转写（含时间戳）" in raw_text
    assert "[00:00:00]" in raw_text
    assert "我们今天呢主要讨论一下项目的进度问题" in raw_text


def test_export_all_formats_and_docx_readback(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    out_dir = tmp_path / "out"
    files = export_result(
        result,
        out_dir=out_dir,
        base_name="项目会议",
        formats=["txt", "md", "docx", "srt"],
        title="项目会议",
    )

    assert sorted(path.suffix for path in files) == [".docx", ".md", ".srt", ".txt", ".txt"]
    md_path = next(path for path in files if path.suffix == ".md")
    md_text = md_path.read_text(encoding="utf-8")
    assert md_text.startswith("# 项目会议")
    assert "> 源音频：会议 录音.mp3" in md_text
    assert "前端大约完成了百分之七十" in md_text

    # docx 读回验证：标题段落 + 正文关键文字
    docx_path = next(path for path in files if path.suffix == ".docx")
    document = Document(str(docx_path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    joined = "\n".join(paragraphs)
    assert paragraphs[0] == "项目会议"
    assert "前端大约完成了百分之七十" in joined
    assert "源音频：会议 录音.mp3" in joined
    # 标题用 Word 标题样式
    assert "Heading" in document.paragraphs[0].style.name


def test_export_without_refined_uses_transcript_and_skips_raw(tmp_path, sample_transcript):
    result = _make_result(sample_transcript, refined_text=None)
    out_dir = tmp_path / "out"
    files = export_result(
        result,
        out_dir=out_dir,
        base_name="未整理",
        formats=["txt", "md", "docx"],
        include_raw=True,  # 没有整理稿时不应生成 _原始转写
    )

    assert len(files) == 3
    assert all("_原始转写" not in path.name for path in files)
    txt_text = files[0].read_text(encoding="utf-8")
    assert "内容：原始转写" in txt_text
    assert "我们今天呢主要讨论一下项目的进度问题" in txt_text
    # 未整理时原始正文不含时间戳前缀
    assert "[00:00:00]" not in txt_text


# ---------------------------------------------------------------------------
# 重名、非法格式、目录
# ---------------------------------------------------------------------------
def test_repeated_export_never_overwrites(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    out_dir = tmp_path / "out"

    first = export_result(result, out_dir=out_dir, base_name="报告", formats=["txt"])
    first_text = first[0].read_text(encoding="utf-8")

    second = export_result(result, out_dir=out_dir, base_name="报告", formats=["txt"])
    third = export_result(result, out_dir=out_dir, base_name="报告", formats=["txt"])

    assert first[0].name == "报告.txt"
    assert second[0].name == "报告 (2).txt"
    assert third[0].name == "报告 (3).txt"
    # 旧文件内容保持原样，没有被覆盖
    assert first[0].read_text(encoding="utf-8") == first_text
    assert first[0].stat().st_size > 0


def test_invalid_formats_fallback_to_txt(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    files = export_result(
        result,
        out_dir=tmp_path / "out",
        base_name="兜底",
        formats=["pdf", "", "TXT", "md", None],  # type: ignore[list-item]
        include_raw=False,
    )
    assert _names(files) == ["兜底.txt", "兜底.md"]


def test_empty_formats_fallback_to_txt(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    files = export_result(
        result,
        out_dir=tmp_path / "out",
        base_name="空",
        formats=[],
        include_raw=False,
    )
    assert len(files) == 1
    assert files[0].suffix == ".txt"


def test_out_dir_created_recursively(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    out_dir = tmp_path / "a" / "b" / "c"
    files = export_result(result, out_dir=out_dir, base_name="嵌套", formats=["txt"])
    assert out_dir.is_dir()
    assert files[0].parent == out_dir


def test_sanitize_weird_base_name(tmp_path, sample_transcript):
    result = _make_result(sample_transcript)
    files = export_result(
        result,
        out_dir=tmp_path / "out",
        base_name='<>:"/\\|?*会议',
        formats=["txt"],
    )
    name = files[0].name
    assert name.endswith(".txt")
    for char in '<>:"/\\|?*\x00':
        assert char not in name


# ---------------------------------------------------------------------------
# SRT
# ---------------------------------------------------------------------------
def test_srt_timestamps_skip_empty_segments_and_keep_sequence(tmp_path):
    transcript = Transcript(
        segments=[
            Segment(0.0, 1.0, "第一句"),
            Segment(1.0, 2.0, "   "),  # 空白片段应跳过
            Segment(2.0, 65.25, "第二句"),
        ],
        language="zh",
        duration=65.25,
    )
    result = _make_result(transcript, refined_text=None, audio_name="字幕.mp3")
    files = export_result(
        result, out_dir=tmp_path / "out", base_name="字幕", formats=["srt"]
    )
    assert len(files) == 1
    srt = files[0].read_text(encoding="utf-8")

    assert "1\n00:00:00,000 --> 00:00:01,000\n第一句" in srt
    assert "2\n00:00:02,000 --> 00:01:05,250\n第二句" in srt
    # 空白段被跳过，序号连续且没有第 3 条
    assert "\n3\n" not in srt
    assert "   " not in srt


def test_srt_without_segments_is_skipped(tmp_path):
    result = _make_result(Transcript(segments=[], duration=0.0), refined_text=None)
    files = export_result(
        result, out_dir=tmp_path / "out", base_name="空转写", formats=["srt"]
    )
    assert files == []


# ---------------------------------------------------------------------------
# 中文长文本与异常
# ---------------------------------------------------------------------------
def test_long_chinese_text_is_fully_written(tmp_path):
    unit = "这是一段很长的中文转写内容，用于验证导出不会截断。"
    long_text = unit * 400  # 约 9200 字
    transcript = Transcript(
        segments=[Segment(0.0, 3600.0, long_text)], language="zh", duration=3600.0
    )
    result = _make_result(transcript, refined_text=long_text)
    files = export_result(
        result, out_dir=tmp_path / "out", base_name="长文本", formats=["txt", "docx"]
    )

    txt_text = files[0].read_text(encoding="utf-8")
    assert txt_text.count(unit) == 400
    assert "音频时长：1小时00分00秒" in txt_text
    # UTF-8 编码写入
    assert files[0].read_bytes().decode("utf-8") == txt_text

    document = Document(str(files[1]))
    docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert docx_text.count(unit) == 400


def test_write_error_wrapped_as_voxerror(tmp_path, sample_transcript, monkeypatch):
    result = _make_result(sample_transcript)

    def _boom(self, *args, **kwargs):  # noqa: ANN001
        raise OSError("磁盘只读")

    monkeypatch.setattr(Path, "write_text", _boom)
    with pytest.raises(VoxError) as excinfo:
        export_result(
            result,
            out_dir=tmp_path / "out",
            base_name="失败",
            formats=["txt"],
        )
    assert "导出" in str(excinfo.value)
    assert excinfo.value.user_hint  # 必须带中文提示


def test_metadata_uses_last_segment_end_when_duration_missing(tmp_path):
    transcript = Transcript(
        segments=[Segment(0.0, 30.0, "前半段"), Segment(30.0, 95.0, "后半段")],
        language="zh",
        duration=0.0,  # 部分来源不提供总时长
    )
    result = _make_result(transcript, refined_text=None)
    files = export_result(result, out_dir=tmp_path / "out", base_name="时长", formats=["txt"])
    text = files[0].read_text(encoding="utf-8")
    assert "音频时长：1分35秒" in text
